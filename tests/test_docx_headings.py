"""A heading must not be merged into the paragraph below it.

Measured failure: a "ДОГОВОР ПОСТАВКИ" Heading 1 followed by a Normal paragraph was treated
as one sentence, so the heading's own translation was lost and the body translation was
written into the heading's runs - the customer's document came back with no title and a
paragraph in Heading 1 style.
"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from app import bootstrap  # noqa: E402

bootstrap.prepare()

failures = 0


def check(ok, label):
    global failures
    print(f"  [{'OK  ' if ok else 'FAIL'}] {label}")
    if not ok:
        failures += 1


print("=" * 74)
print("Word：标题不能被并进正文")
print("=" * 74)

try:
    from docx import Document
except ModuleNotFoundError:
    print("  跳过：这台机器没装 python-docx")
    sys.exit(0)

from app.config import AppConfig  # noqa: E402
from app.docx_io import translate_docx  # noqa: E402

# A fake translator so the test does not need a model.
def fake(text: str) -> str:
    return "[译]" + text.strip()

work = Path(tempfile.mkdtemp(prefix="ruzh-docx-"))
src = work / "in.docx"
doc = Document()
doc.add_heading("ДОГОВОР ПОСТАВКИ", level=1)
doc.add_paragraph("Общая сумма договора.")
doc.add_paragraph("Вторая строка.")
doc.save(src)

dst = work / "out.docx"
cfg = AppConfig()
cfg.source_lang, cfg.target_lang = "ru", "zh"
translate_docx(cfg, cfg.resolved_models_root(), src, dst, translator=fake)

out = Document(str(dst))
texts = [p.text for p in out.paragraphs]
styles = [p.style.name for p in out.paragraphs]

check(len(out.paragraphs) >= 3, f"段数没被吃掉：{len(out.paragraphs)}")
check("ДОГОВОР ПОСТАВКИ" in texts[0], f"标题在第一段且单独成段：{texts[0]!r}")
check(styles[0].lower().startswith("heading"), f"标题样式还在：{styles[0]}")
check("Общая" not in texts[0], "正文没有被写进标题")

print()
print("=" * 74)
print(f"失败 {failures} 项")
print("=" * 74)
sys.exit(1 if failures else 0)
